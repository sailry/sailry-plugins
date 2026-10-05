"""Copy and adapt this starter; use the full library API for the requested document."""
import argparse
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('output', type=Path)
output = parser.parse_args().output
if output.exists():
    parser.error('destination already exists; choose a new path')
output.parent.mkdir(parents=True, exist_ok=True)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

presentation = Presentation()
presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
heading = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12), Inches(.8))
run = heading.text_frame.paragraphs[0].add_run()
run.text, run.font.size = 'Quarterly results', Pt(32)
data = CategoryChartData()
data.categories = ['Q1', 'Q2', 'Q3']
data.add_series('Revenue', [100, 140, 180])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(10), Inches(4.5), data)
slide.notes_slide.notes_text_frame.text = 'Explain the growth trend.'
presentation.save(output)
assert len(Presentation(output).slides) == 1
