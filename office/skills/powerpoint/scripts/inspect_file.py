"""Inspect saved content; this does not replace layout or feature-specific validation."""
import argparse
import json
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('source', type=Path)
source = parser.parse_args().source

from pptx import Presentation

presentation = Presentation(source)
print(json.dumps([{
    'slide': index + 1,
    'text': [shape.text for shape in slide.shapes if shape.has_text_frame],
    'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else '',
} for index, slide in enumerate(presentation.slides)], ensure_ascii=False, indent=2))
