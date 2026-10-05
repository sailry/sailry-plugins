"""Copy and adapt this starter; use the full library API for the requested document."""
import argparse
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('output', type=Path)
output = parser.parse_args().output
if output.exists():
    parser.error('destination already exists; choose a new path')
output.parent.mkdir(parents=True, exist_ok=True)

from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

document = Document()
section = document.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Mm(25)
section.left_margin = section.right_margin = Mm(28)
normal = document.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
document.add_heading('Report', 0)
document.add_heading('Overview', 1)
document.add_paragraph('Replace this paragraph with the requested content.')
table = document.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Item'
table.rows[0].cells[1].text = 'Details'
row = table.add_row()
row.cells[0].text = 'Summary'
row.cells[1].text = 'Add the supporting information.'
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
section.footer.paragraphs[0]._p.append(field)
document.save(output)
assert Document(output).tables[0].rows[1].cells[0].text == 'Summary'
