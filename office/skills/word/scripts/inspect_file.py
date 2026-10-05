"""Inspect saved content; this does not replace layout or feature-specific validation."""
import argparse
import json
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('source', type=Path)
source = parser.parse_args().source

from docx import Document

document = Document(source)
print(json.dumps({
    'paragraphs': [p.text for p in document.paragraphs],
    'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
    'headers': [[p.text for p in section.header.paragraphs] for section in document.sections],
    'footers': [[p.text for p in section.footer.paragraphs] for section in document.sections],
}, ensure_ascii=False, indent=2))
