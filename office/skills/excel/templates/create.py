"""Copy and adapt this starter; use the full library API for the requested document."""
import argparse
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('output', type=Path)
output = parser.parse_args().output
if output.exists():
    parser.error('destination already exists; choose a new path')
output.parent.mkdir(parents=True, exist_ok=True)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference

workbook = Workbook()
sheet = workbook.active
sheet.title = 'Budget'
for row in [('Item', 'Amount'), ('Design', 120), ('Build', 240), ('Total', '=SUM(B2:B3)')]:
    sheet.append(row)
for cell in sheet[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='245A81')
sheet.column_dimensions['A'].width = 24
sheet.column_dimensions['B'].width = 18
sheet.freeze_panes = 'A2'
sheet.auto_filter.ref = 'A1:B3'
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
sheet.add_chart(chart, 'D2')
sheet.print_options.horizontalCentered = True
sheet.print_area = 'A1:L18'
workbook.save(output)
assert load_workbook(output)['Budget']['B4'].value == '=SUM(B2:B3)'
# openpyxl preserves formulas but does not calculate their cached results.
